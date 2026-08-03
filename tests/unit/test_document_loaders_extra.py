from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument
from openpyxl import Workbook
from pypdf import PdfWriter

from ragzen.loaders.documents import UniversalDocumentLoader


def test_docx_xlsx_and_pdf_loaders(tmp_path: Path) -> None:
    loader = UniversalDocumentLoader()

    docx_path = tmp_path / "policy.docx"
    document = DocxDocument()
    document.add_paragraph("Document policy text")
    document.save(docx_path)
    assert "Document policy" in loader.load(docx_path)[0].content

    xlsx_path = tmp_path / "policy.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["Policy", "30 days"])
    workbook.save(xlsx_path)
    workbook.close()
    assert "30 days" in loader.load(xlsx_path)[0].content

    pdf_path = tmp_path / "empty.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with pdf_path.open("wb") as output:
        writer.write(output)
    loaded = loader.load(pdf_path)[0]
    assert loaded.page_count == 1
